"""
build_canonical_export.py
=========================
Builds the format-independent canonical export of the final LOVD table and the review
workbook for the IDbases curator. Offline; reads committed pipeline outputs only.

canonical_export.tsv   one row per final patient-variant pair (lovd_flat_with_patients_final.tsv),
                       except records that describe no change (review_reason "no change described
                       (c.=)"). Values copied from the final table are byte-equal to it.
review_workbook_for_mauno.xlsx
                       README, Review_tier (every Review pair), Suspected_source_errors
                       (docs/agents/source_errors_for_mauno.csv) and Counts.

Allele, phase and zygosity come from the pair's own source change (source_change_id: gene |
accession | allele | g. | c.) and the source units of the individual (audit_variants.tsv):
    allele      the source allele number of the change (0 = not assigned in the source)
    zygosity    homozygous when the source unit is the same change on alleles 1 and 2 ("1+2");
                hemizygous for a chromosome X gene in a male; heterozygous when the individual
                carries a different change on the other allele; unknown otherwise
    phase       cis with the other changes of the same source allele (bracketed in-cis
                notation), trans with the changes on the other allele, unknown for allele 0
Gene symbol: Output/alias.csv for renamed IDbases symbols, else the gene name of the gene's
MANE Select transcript (MANE v1.5 GFF), else the IDbases symbol (checked against the MANE v1.5
gene names). metadata_flag marks a sex field that is not M, F or empty (multi-individual strings).
For pairs whose record is NMD_predicted (Step9_NMD) r_hgvs and p_hgvs are r.0? and p.0?, the
Mutalyzer r. and p. move into alt_descriptions, and nmd_rule_applied says so; every other copied
value is byte-equal to the final table. The workbook's Suspected_source_errors sheet leaves out
the source changes whose kept record is NMD_predicted, since their protein is not described.

Usage:
    python Step11_Export/build_canonical_export.py [--out-dir Output/Step11_Export]
"""

import argparse
import csv
import gzip
import re
from collections import Counter, defaultdict
from pathlib import Path

SCRIPTS = Path(__file__).resolve().parents[1]
REPO = SCRIPTS.parents[1]
OUT_ROOT = SCRIPTS.parent / "Output"
MANE_GFF = REPO / "01_Reference_Genomes" / "GRCh38_MANE" / "MANE.GRCh38.v1.5.ensembl_genomic.gff.gz"
NO_CHANGE = "no change described (c.=)"
EXPORT_COLS = [
    "individual_id", "gene_symbol", "gene_symbol_source", "idbase_gene", "accession", "source_change_id",
    "allele", "zygosity", "phase", "phase_partners",
    "assembly", "chromosome", "genomic_hgvs", "transcript", "c_hgvs", "r_hgvs", "p_hgvs",
    "variant_class", "vario_dna", "vario_protein", "nmd_call", "alt_descriptions",
    "published_as", "source_track", "confidence_tier", "review_reason", "source_error", "source_error_category",
    "source_error_severity",
    "sex", "geographic_origin", "age", "individual_remarks", "phenotype", "phenotype_additional",
    "genome_remarks", "pubmed_id", "reference_authors", "reference_title", "reference_location",
    "screening", "genetic_origin", "effect", "metadata_flag", "nmd_rule_applied",
]
NMD_RULE = "predicted NMD: stop more than 50 nt before the last coding exon junction; r. and p. set to r.0? and p.0?"
COPIED = {
    "individual_id": "Individual/ID", "chromosome": "chromosome", "genomic_hgvs": "VariantOnGenome/DNA",
    "transcript": "transcriptid", "c_hgvs": "VariantOnTranscript/DNA", "r_hgvs": "VariantOnTranscript/RNA",
    "p_hgvs": "VariantOnTranscript/Protein", "variant_class": "_variant_class", "vario_dna": "VariO/DNA",
    "vario_protein": "VariO/Protein", "alt_descriptions": "alt_descriptions", "source_track": "_source_track",
    "confidence_tier": "confidence_tier", "review_reason": "review_reason", "sex": "Individual/Gender",
    "geographic_origin": "Individual/Origin/Geographic", "age": "Individual/Age",
    "individual_remarks": "Individual/Remarks", "phenotype": "Phenotype/Disease",
    "phenotype_additional": "Phenotype/Additional", "genome_remarks": "VariantOnGenome/Remarks",
    "pubmed_id": "Reference/PubMed_ID", "reference_authors": "Reference/Authors",
    "reference_title": "Reference/Title", "reference_location": "Reference/Location",
    "screening": "Screening/Technique",
}
RE_C_GROUP = re.compile(r"c\.\s*\[([^\]]*)\]")
RE_C = re.compile(r"c\.(?!\[)[^,\];\s\[]+")


def c_tokens(text):
    out = ["c." + re.sub(r"^c\.", "", x.strip()) for m in RE_C_GROUP.finditer(text) for x in re.split(r";|\s\+\s", m.group(1)) if x.strip()]
    out += RE_C.findall(RE_C_GROUP.sub("", text))
    return list(dict.fromkeys(out))


def read_tsv(path):
    with open(path, encoding="utf-8") as f:
        return list(csv.DictReader(f, delimiter="\t"))


def mane_gene_names(gff):
    out = {}
    with gzip.open(gff, "rt", encoding="utf-8") as f:
        for line in f:
            if "\ttranscript\t" in line and "tag=MANE_Select" in line:
                a = dict(x.split("=", 1) for x in line.rstrip("\n").split("\t")[8].split(";") if "=" in x)
                nm = re.search(r"RefSeq:(NM_\d+\.\d+)", a.get("Dbxref", ""))
                if nm:
                    out[nm.group(1)] = a["gene_name"]
    return out


def unit_alleles(label):
    return set(label.split("#", 1)[0].split("+"))


class Units:
    def __init__(self, audit_variants):
        self.by_acc = defaultdict(list)
        for r in audit_variants:
            self.by_acc[(r["gene"].strip().upper(), r["accession"].strip())].append(r)

    def of(self, sid):
        gene, acc, al, g, c = sid.split("|", 4)
        cands = [u for u in self.by_acc.get((gene, acc), []) if al in unit_alleles(u["allele"])]
        hit = [u for u in cands if c != "-" and c in c_tokens(u["source_systematic_name"])]
        return (hit or cands or [None])[0]


def zygosity_phase(sid, unit, units, chrom, sex):
    gene, acc, al, g, c = sid.split("|", 4)
    own = unit["source_systematic_name"] if unit else ""
    cis = [x for x in c_tokens(own) if x != c] if unit else []
    others = [u for u in units.by_acc.get((gene, acc), []) if u is not unit]
    other_allele = "2" if al == "1" else ("1" if al == "2" else "")
    trans = [x for u in others if other_allele and other_allele in unit_alleles(u["allele"]) for x in c_tokens(u["source_systematic_name"])]
    if unit and unit_alleles(unit["allele"]) >= {"1", "2"}:
        zyg = "homozygous"
    elif chrom == "X" and sex.strip() == "M":
        zyg = "hemizygous"
    elif trans:
        zyg = "heterozygous"
    else:
        zyg = "unknown"
    if al == "0" or zyg == "homozygous":
        phase = "unknown" if al == "0" else "both alleles"
    else:
        phase = "+".join(x for x, ok in (("cis", cis), ("trans", trans)) if ok) or "unknown"
    partners = "; ".join([f"cis {x}" for x in cis] + [f"trans {x}" for x in dict.fromkeys(trans)])
    return zyg, phase, partners


def build(out_dir):
    final = read_tsv(OUT_ROOT / "Step8_Rescue" / "lovd_flat_with_patients_final.tsv")
    pairs = {(r["individual_id"], r["dedup_key"]): r for r in read_tsv(OUT_ROOT / "Step8_Rescue" / "resolution_pairs.tsv")}
    units = Units(read_tsv(OUT_ROOT / "Step10_Audit" / "audit_variants.tsv"))
    nmd = {(r["_gene"], r["_dedup_key"]): r["nmd_call"] for r in read_tsv(OUT_ROOT / "Step9_NMD" / "nmd_annotation.tsv")}
    alias = {r["gene"].strip().upper(): r["alias"].strip() for r in csv.DictReader(open(OUT_ROOT / "alias.csv", encoding="utf-8"))}
    mane_name = mane_gene_names(MANE_GFF)
    mane_genes = set(mane_name.values())
    gene_name = {}
    for f in final:
        if f["transcriptid"] in mane_name:
            gene_name.setdefault(f["_gene"].strip().upper(), mane_name[f["transcriptid"]])
    se = {r["source_change_id"]: r for r in read_tsv(SCRIPTS / "refs" / "source_errors.tsv")}
    rows, checks = [], Counter()
    for f in final:
        if f["review_reason"] == NO_CHANGE:
            checks["excluded_no_change"] += 1
            continue
        pr = pairs.get((f["Individual/ID"], f["_dedup_key"]), {})
        sids = [s.strip() for s in pr.get("source_change_ids", "").split(" ; ") if s.strip()]
        sid = sids[0] if sids else ""
        unit = units.of(sid) if sid else None
        gene = f["_gene"].strip().upper()
        if gene in alias:
            sym, src = alias[gene], "alias.csv"
        elif gene in gene_name:
            sym, src = gene_name[gene], "MANE v1.5 transcript"
        elif f["_gene"] in mane_genes:
            sym, src = f["_gene"], "MANE v1.5 gene name"
        else:
            sym, src = f["_gene"], "IDbases"
        zyg, phase, partners = zygosity_phase(sid, unit, units, f["chromosome"], f["Individual/Gender"]) if sid else ("unknown", "unknown", "")
        e = se.get(sid, {})
        row = {k: f[v] for k, v in COPIED.items()}
        row.update({
            "gene_symbol": sym, "gene_symbol_source": src, "idbase_gene": f["_gene"],
            "accession": f["Individual/ID"].split("_", 1)[1], "source_change_id": sid,
            "allele": sid.split("|")[2] if sid else "", "zygosity": zyg, "phase": phase, "phase_partners": partners,
            "assembly": "GRCh38", "nmd_call": nmd.get((f["_gene"], f["_dedup_key"]), ""),
            "published_as": unit["source_systematic_name"] if unit else "",
            "source_error": "candidate" if e else "", "source_error_category": e.get("category", ""),
            "source_error_severity": e.get("severity", ""),
            "genetic_origin": "", "effect": "",
            "metadata_flag": "" if f["Individual/Gender"].strip() in ("", "M", "F") else "sex field garbled (multi-individual string)",
            "nmd_rule_applied": "",
        })
        if row["nmd_call"] == "NMD_predicted":
            moved = " ".join(x for x in (f["VariantOnTranscript/RNA"], f["VariantOnTranscript/Protein"]) if x)
            row["alt_descriptions"] = " | ".join(x for x in (f["alt_descriptions"], f"Mutalyzer {moved}" if moved else "") if x)
            row["r_hgvs"], row["p_hgvs"], row["nmd_rule_applied"] = "r.0?", "p.0?", NMD_RULE
            checks["nmd_rule_rows"] += 1
        rows.append(row)
        checks["no_sid"] += not sid
        checks["no_unit"] += unit is None
        nmd_cols = {"r_hgvs", "p_hgvs", "alt_descriptions"} if row["nmd_rule_applied"] else set()
        checks["copy_mismatch"] += sum(row[k] != f[v] for k, v in COPIED.items() if k not in nmd_cols)
        checks["nmd_alt_not_prefix"] += bool(nmd_cols) and not row["alt_descriptions"].startswith(f["alt_descriptions"])
    out_dir.mkdir(parents=True, exist_ok=True)
    with open(out_dir / "canonical_export.tsv", "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=EXPORT_COLS, delimiter="\t", lineterminator="\n")
        w.writeheader()
        w.writerows(rows)
    return final, rows, checks, units


def plain_reason(f, log):
    r = f["review_reason"]
    if r.startswith("genomic record"):
        return "Our genomic position for this variant is outside the gene's MANE transcript or deep in an intron, while the published c. is exonic: our mapping is probably wrong. " + r
    if r.startswith("record keeps the literal"):
        return "Our record uses the published c. literally, but this gene's published numbering follows an older scheme. " + r
    if r == NO_CHANGE:
        return "The published notation describes no change (reference and new base are the same)."
    if r:
        return r
    lg = log.get((f["Individual/ID"], f["_dedup_key"]), {})
    alts = lg.get("dropped", "")
    return ("Different pipeline routes gave different positions for this change and none could be confirmed against the published protein; "
            f"kept by {lg.get('deciding_signal', 'track resolution')}" + (f"; alternatives: {alts}" if alts else "") + ".")


def workbook(out_dir, final, rows, units):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    head_fill, head_font = PatternFill("solid", start_color="305496"), Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    ask_fill = PatternFill("solid", start_color="FFF2CC")
    pairs = {(r["individual_id"], r["dedup_key"]): r for r in read_tsv(OUT_ROOT / "Step8_Rescue" / "resolution_pairs.tsv")}
    log = {}
    for r in read_tsv(OUT_ROOT / "Step8_Rescue" / "resolution_log.tsv"):
        log.setdefault((r["individual_id"], r["kept"]), r)
    wb = Workbook()

    def sheet(ws, header, data, ask_col=None):
        ws.append(header)
        for c in ws[1]:
            c.fill, c.font = head_fill, head_font
        for d in data:
            ws.append(d)
        for i, h in enumerate(header, 1):
            width = max([len(str(h))] + [len(str(d[i - 1])) for d in data[:500]])
            ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, width + 2))
            if h == ask_col:
                for cell in ws[get_column_letter(i)][1:]:
                    cell.fill = ask_fill
        ws.freeze_panes = "A2"

    review = []
    for f in final:
        if f["confidence_tier"] != "Review":
            continue
        pr = pairs.get((f["Individual/ID"], f["_dedup_key"]), {})
        sid = pr.get("source_change_ids", "").split(" ; ")[0]
        u = units.of(sid) if sid else None
        review.append([f["_gene"], f["Individual/ID"].split("_", 1)[1], sid.split("|")[2] if sid else "",
                       u["src_g"] if u else "", u["src_c"] if u else "", u["src_p"] if u else "",
                       f["transcriptid"], f["VariantOnTranscript/DNA"], f["VariantOnTranscript/Protein"],
                       plain_reason(f, log), ""])
    src_all = list(csv.reader(open(REPO / "docs" / "agents" / "source_errors_for_mauno.csv", encoding="utf-8-sig")))
    head = src_all[0]
    by_change = defaultdict(set)
    for r in rows:
        if r["source_change_id"]:
            g, acc, _, _, c = r["source_change_id"].split("|", 4)
            by_change[(g, acc, c)].add(r["nmd_call"])
            by_change[(r["gene_symbol"].upper(), acc, c)].add(r["nmd_call"])
    src_err, dropped = [head], []
    for x in src_all[1:]:
        d = dict(zip(head, x))
        calls = set().union(*[by_change.get((d["gene"].strip().upper(), a.strip(), d["source_c"].strip()), set()) for a in d["accessions"].split(";")])
        (dropped if calls == {"NMD_predicted"} else src_err).append(x)
    readme = [
        ["Review workbook for the IDbases migration to LOVD 3.0 (GRCh38, MANE Select, HGVS)."],
        ["Review_tier: patient-variant pairs we will not submit as confirmed. Each row says why (column review_reason)."],
        ["  Please choose keep (submit as is), correct (give the right notation) or drop, in Mauno_decision."],
        ["Suspected_source_errors: published notations that contradict themselves (c. vs p., g. vs c.) or describe no change."],
        ["  Please confirm or correct each in Mauno_decision; why_flagged explains what looks inconsistent."],
        ["Counts: the number of rows in each sheet, computed from the data."],
        ["Our c./p. are on the MANE Select transcript; 'source' columns are the notation as published in IDbases."],
        ["Nothing in this workbook has been submitted; decisions here feed the final export."],
    ]
    ws = wb.active
    ws.title = "README"
    for line in readme:
        ws.append(line)
    ws["A1"].font = Font(name="Calibri", bold=True, size=14, color="1F3864")
    ws.column_dimensions["A"].width = 120
    sheet(wb.create_sheet("Review_tier"), ["gene", "accession", "allele", "source_g", "source_c", "source_p", "our_transcript",
                                           "our_c", "our_p", "review_reason", "Mauno_decision"], review, "Mauno_decision")
    sheet(wb.create_sheet("Suspected_source_errors"), src_err[0] + ["Mauno_decision"], [r + [""] for r in src_err[1:]], "Mauno_decision")
    counts = [["Review_tier pairs", len(review)]]
    kinds = Counter("genomic position" if f["review_reason"].startswith("genomic record") else
                    "literal numbering" if f["review_reason"].startswith("record keeps the literal") else
                    "no change" if f["review_reason"] == NO_CHANGE else "track conflict"
                    for f in final if f["confidence_tier"] == "Review")
    counts += [[f"  of which {k}", v] for k, v in sorted(kinds.items())]
    counts += [["Suspected_source_errors rows", len(src_err) - 1],
               ["  left out: kept record predicted NMD (protein not described)", len(dropped)],
               ["Canonical export rows", len(rows)]]
    sheet(wb.create_sheet("Counts"), ["item", "n"], counts)
    for name in ("Review_tier", "Suspected_source_errors"):
        for row in wb[name].iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=False, vertical="top")
    wb.save(out_dir / "review_workbook_for_mauno.xlsx")
    return review, counts, [dict(zip(head, x)) for x in dropped]


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--out-dir", type=Path, default=OUT_ROOT / "Step11_Export")
    args = ap.parse_args()
    final, rows, checks, units = build(args.out_dir)
    review, counts, dropped = workbook(args.out_dir, final, rows, units)
    print(f"final pairs {len(final)}; excluded (no change) {checks['excluded_no_change']}; export rows {len(rows)}")
    print(f"rows without source_change_id {checks['no_sid']}; without source unit {checks['no_unit']}; copied-value mismatches {checks['copy_mismatch']}")
    print(f"NMD rule applied to {checks['nmd_rule_rows']} pairs; alt_descriptions not extending the original: {checks['nmd_alt_not_prefix']}")
    print("source-error rows left out (kept record NMD_predicted):", [(d["gene"], d["accessions"], d["source_c"]) for d in dropped])
    print("tiers:", dict(Counter(r["confidence_tier"] for r in rows)), "tracks:", dict(Counter(r["source_track"] for r in rows)))
    print("zygosity:", dict(Counter(r["zygosity"] for r in rows)), "phase:", dict(Counter(r["phase"] for r in rows)))
    dup = Counter((r["individual_id"], r["c_hgvs"]) for r in rows)
    print("duplicate (individual_id, c.) keys:", sum(1 for v in dup.values() if v > 1))
    print("symbols differing from IDbases:", sorted({(r["idbase_gene"], r["gene_symbol"], r["gene_symbol_source"]) for r in rows if r["gene_symbol"] != r["idbase_gene"]}))
    print("workbook:", counts)


if __name__ == "__main__":
    main()
